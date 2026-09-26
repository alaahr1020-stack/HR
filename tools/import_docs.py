#!/usr/bin/env python3
"""Builds the app content (assets/content/library.json + assets/forms/) from
a folder of HR files.

Usage:
    pip install python-docx pypdf openpyxl
    python tools/import_docs.py "C:\\Users\\Alaa Helal\\Downloads\\لائحة تنظيم العمل و الجزاءات"

Supported inputs:
  .docx .xlsx .md .txt  read directly
  .rtf .doc             converted to .docx with LibreOffice (soffice)
  .pdf                  text PDFs read directly; scanned PDFs go through
                        Tesseract OCR (Arabic) when it is installed
Conversions and OCR results are cached in tools/.cache.

How files are organised in the app:
  * "نماذج" / "كشف" in the path     → نماذج الموارد البشرية
    A Word file holding many forms (tables that start with "نموذج /") is
    split so every form becomes its own item with its own .docx to share.
  * "لائحة نموذجية" in the name      → لوائح تنظيم العمل والجزاءات
    When "X - نسخة المراجعة" and "X" both exist, the review copy (with the
    legal basis of each article) is shown and the clean copy is attached.
  * "الالتزامات" / "العقوبات"        → التزامات صاحب العمل والعقوبات
  * anything else                     → التشريعات والقرارات, grouped by folder
"""

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT_JSON = ROOT / "assets" / "content" / "library.json"
FORMS_DIR = ROOT / "assets" / "forms"
CACHE = Path(__file__).resolve().parent / ".cache"

CATEGORIES = [
    {"id": "laws", "title": "التشريعات والقرارات", "icon": "gavel",
     "description": "قانون العمل وقراراته الوزارية والقوانين المرتبطة"},
    {"id": "work_rules", "title": "لوائح تنظيم العمل والجزاءات", "icon": "rule",
     "description": "لوائح نموذجية بجداول الجزاءات وسند كل مادة"},
    {"id": "obligations", "title": "التزامات صاحب العمل والعقوبات", "icon": "warning",
     "description": "كل التزام قانوني ومادته والغرامة المقررة عند مخالفته"},
    {"id": "forms", "title": "نماذج الموارد البشرية", "icon": "forms",
     "description": "نماذج وورد وإكسل جاهزة للحفظ والمشاركة"},
]

# Display names for law folders, in display order; other folders follow
# alphabetically under their own names.
GROUP_NAMES = {
    "قانون العمل": "قانون العمل 14 لسنة 2025 وقراراته",
    "قانون التامينات": "التأمينات الاجتماعية والمعاشات",
    "قانون التامين الصحي الشامل": "التأمين الصحي الشامل",
    "صندوق اعانات الطوارئ": "صندوق إعانات الطوارئ للعمال",
    "قانون ذوي الاعاقة": "حقوق الأشخاص ذوي الإعاقة",
    "قانون المنظمات النقابية": "المنظمات النقابية العمالية",
    "قانون حماية البيانات الشخصية": "حماية البيانات الشخصية",
    "المحكمة الدستورية": "أحكام المحكمة الدستورية العليا",
    "جرائم تقنية المعلومات": "مكافحة جرائم تقنية المعلومات",
    "قانون الطفل": "قانون الطفل",
    "قانون التعليم و تعديلاته": "قانون التعليم وتعديلاته",
    "قانون المنشآت السياحية و الفندقية": "المنشآت الفندقية والسياحية",
    "قانون المهن الطبية": "المسئولية الطبية وسلامة المريض",
    "قانون تنظيم الصناعة": "تنظيم الصناعة",
    "قانون الخدمة العسكرية": "الخدمة العسكرية والوطنية",
}
GROUP_ORDER = list(GROUP_NAMES.values())

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MAX_SECTION = 4000
ARTICLE_RE = re.compile(
    r"^\s*[(\[]?\s*(ال)?(مادة|المادة|باب|الباب|فصل|الفصل|بند|البند|الفرع|القسم|ملحق|الملحق|جدول)\b"
    r"|^\s*(أولا|اولا|ثانيا|ثالثا|رابعا|خامسا|سادسا|سابعا|ثامنا|تاسعا|عاشرا)[ًٌ]?\s*[:\-–]"
    r"|^\s*(إجراء|اجراء|خطوة|الخطوة)\s*(رقم)?\s*[\d٠-٩]+"
    r"|^\s*ديباجة\s*$"
)
CHAPTER_RE = re.compile(r"^\s*(ال)?(باب|فصل|فرع|قسم|ملحق)\b")


# ---------------------------------------------------------------- helpers

def log(msg):
    print(msg, flush=True)


def file_hash(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def fix_mojibake(text: str) -> str:
    """Arabic stored as Windows-1256 but decoded as Windows-1252 turns into
    'æÒÇÑÉ ÇáÚãá'. Re-decode such words."""
    def fix(m):
        w = m.group(0)
        try:
            return w.encode("cp1252").decode("cp1256")
        except UnicodeError:
            try:
                return w.encode("latin-1").decode("cp1256")
            except UnicodeError:
                return w
    return re.sub(r"[\u00c0-\u00ff\u0152\u0153\u0160\u0161\u0178\u0192\u02c6\u02dc]{2,}", fix, text)


def clean(text: str) -> str:
    text = fix_mojibake(text)
    text = text.replace("\u00a0", " ").replace("\r", "").replace("\u200f", "").replace("\u200e", "")
    text = re.sub(r"(?<=\S)ـ+|ـ+(?=\S)", "", text)  # tatweel
    text = re.sub(r"[ \t]+", " ", to_latin_digits(text))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def to_latin_digits(s: str) -> str:
    return s.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))


def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def convert_to_docx(path: Path) -> Path:
    out = CACHE / "docx" / f"{file_hash(path)}.docx"
    if out.exists():
        return out
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("يلزم تثبيت LibreOffice لتحويل ملفات rtf/doc")
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"in{path.suffix.lower()}"
        shutil.copy2(path, src)
        profile = Path(tmp, "profile").as_uri()
        subprocess.run([soffice, "--headless", f"-env:UserInstallation={profile}",
                        "--convert-to", "docx", "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=1800)
        shutil.move(str(Path(tmp) / "in.docx"), out)
    return out


def ocr_pdf(path: Path) -> str:
    out = CACHE / "ocr" / f"{file_hash(path)}.txt"
    if out.exists():
        return out.read_text(encoding="utf-8")
    if not (shutil.which("tesseract") and shutil.which("pdftoppm")):
        return ""
    out.parent.mkdir(parents=True, exist_ok=True)
    pages = []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-r", "300", "-png", str(path), f"{tmp}/p"], check=True)
        for img in sorted(Path(tmp).glob("p-*.png")):
            r = subprocess.run(["tesseract", str(img), "stdout", "-l", "ara"],
                               capture_output=True, text=True)
            pages.append(r.stdout)
    text = "\n".join(pages)
    out.write_text(text, encoding="utf-8")
    return text


# ------------------------------------------------------------- readers
# Every reader returns a list of blocks: (kind, text) where kind is
# "h" (heading), "p" (paragraph) or "t" (table rows joined by newlines).

def _para_is_heading(p) -> bool:
    try:
        style = (p.style.name or "").lower() if p.style is not None else ""
    except Exception:
        style = ""
    if style.startswith("heading") or style == "title":
        return True
    runs = [r for r in p.runs if r.text.strip()]
    if not runs or not all(r.bold for r in runs):
        return False
    sizes = [r.font.size.pt for r in runs if r.font.size]
    return bool(sizes) and max(sizes) >= 12.5


def _table_rows(tbl):
    rows = []
    for r in tbl.rows:
        cells = []
        for c in r.cells:
            ct = " ".join(c.text.split())
            if ct and (not cells or cells[-1] != ct):  # merged cells repeat
                cells.append(ct)
        if cells:
            rows.append(cells)
    return rows


def docx_blocks(d):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks = []
    for child in d.element.body.iterchildren():
        if child.tag == W_NS + "p":
            p = Paragraph(child, d)
            t = clean(p.text)
            if t:
                blocks.append(("h" if _para_is_heading(p) else "p", t))
        elif child.tag == W_NS + "tbl":
            tbl = Table(child, d)
            rows = _table_rows(tbl)
            # Layout tables (text put inside 1–2 wide columns) → paragraphs.
            if len(tbl.columns) <= 2 and sum(len(" ".join(r)) for r in rows) > 600:
                for c in tbl._cells:
                    for p in c.paragraphs:
                        t = clean(p.text)
                        if t:
                            blocks.append(("p", t))
            elif rows:
                blocks.append(("t", clean("\n".join(" | ".join(r) for r in rows))))
    # A cell's paragraphs repeat for merged cells; drop consecutive duplicates.
    out = []
    for b in blocks:
        if not out or out[-1] != b:
            out.append(b)
    return out


def read_pdf(path: Path):
    from pypdf import PdfReader
    text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
    ocr = False
    if len(text.strip()) < 200:
        text, ocr = ocr_pdf(path), True
    lines = [clean(l) for l in text.splitlines()]
    return [("p", l) for l in lines if len(l) > 1], ocr


def read_xlsx(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    blocks = []
    for ws in wb.worksheets:
        blocks.append(("h", ws.title))
        for row in ws.iter_rows(values_only=True):
            cells = [clean(str(v)) for v in row if v is not None and str(v).strip()]
            if cells:
                blocks.append(("p", " | ".join(cells)))
    return blocks


def read_text(path: Path):
    blocks = []
    for l in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        l = l.strip()
        if not l or re.fullmatch(r"[-|: ]+", l) or l == "---":
            continue
        if l.startswith("#"):
            blocks.append(("h", clean(l.lstrip("# ").replace("**", ""))))
        elif l.startswith("|"):
            cells = [c.strip().replace("**", "") for c in l.strip("|").split("|")]
            blocks.append(("t", clean(" | ".join(c for c in cells if c))))
        else:
            blocks.append(("p", clean(l.replace("**", ""))))
    return blocks


# ------------------------------------------------------------- sections

END_OF_SENTENCE = re.compile(r"[.:：؛;!؟?)\]»\"]\s*$")
LIST_START = re.compile(r"^\s*([\d٠-٩]+\s*[-–.)]|[-–•*]|\(|[أ-ي]\s*[-–)]\s)")


def join_wrapped(lines):
    """Old RTF/PDF laws break every printed line into its own paragraph.
    Re-join lines that continue the previous sentence."""
    out = []
    for l in (x.strip() for line in lines for x in line.split("\n")):
        if not l:
            continue
        if out and "|" not in l and "|" not in out[-1] and not LIST_START.match(l) \
                and not out[-1].startswith("[") \
                and not END_OF_SENTENCE.search(out[-1]) and len(out[-1]) > 25:
            out[-1] = f"{out[-1]} {l}"
        else:
            out.append(l)
    return out


def to_sections(blocks):
    """Groups blocks into sections. Headings without body (الباب/الفصل) are
    kept as the `chapter` label of the next section."""
    sections, title, body, pending = [], "", [], []

    def flush():
        nonlocal pending
        text = "\n".join(join_wrapped(body)).strip()
        t = re.sub(r"^\(\s*(.*?)\s*\)$", r"\1", title.strip()).rstrip(" :：-–")
        if not text:
            if t:
                pending.append(t)
            return
        first = True
        while text:
            cut = len(text)
            if cut > MAX_SECTION:
                cut = text.rfind("\n", 0, MAX_SECTION)
                cut = cut if cut > MAX_SECTION // 2 else MAX_SECTION
            s = {"title": t, "body": text[:cut].strip()}
            if first and pending:
                s["chapter"] = " › ".join(pending[-2:])
                pending = []
            sections.append(s)
            text, first = text[cut:].strip(), False

    for kind, text in blocks:
        if kind == "t":
            body.append(text)
            continue
        short = len(text) <= 120
        split = re.search(r"[:：\-–]\s*\S", text[:60])
        if ARTICLE_RE.match(text) and split and len(text) - split.end() > 40 \
                and not CHAPTER_RE.match(text):
            # "الخطوة 1: نص طويل…" → title + body
            flush()
            head, rest = re.split(r"\s*[:：\-–]\s*", text, maxsplit=1)
            title, body = head.strip(), [rest]
        elif short and (kind == "h" or ARTICLE_RE.match(text)):
            flush()
            if title and not body and not CHAPTER_RE.match(title) and not CHAPTER_RE.match(text):
                title = f"{title} {text}"  # "مادة (1)" followed by its caption
            else:
                title, body = text, []
        else:
            body.append(text)
    flush()
    return sections


# ------------------------------------------------------- legal documents

TITLE_STOP = re.compile(r"^(باسم الشعب|باسم الأمة|باسم الامة|رئيس الجمهورية$|رئيس مجلس الوزراء$|"
                        r"وزير |وزيرة |مجلس الوزراء$|بعد الاطلاع|بعد الإطلاع|نحن |قرر|المادة|مادة)")


def legal_title(blocks):
    """'وزارة العمل / قرار رقم 48 لسنة 2026 / بشأن …' → one clean title."""
    lines = []
    for _, t in blocks[:12]:
        for l in t.split("\n"):
            l = l.strip(" /")
            if l:
                lines.append(l)
    ministry, parts = None, []
    for l in lines:
        if not parts and l.startswith("وزارة"):
            ministry = l
            continue
        if parts and TITLE_STOP.match(l):
            break
        if not parts and not re.search(r"(قانون|قرار|مرسوم|القانون)", l):
            if len(lines) and re.match(r"^\d", l):  # "27 مارس 1948 - قانون…"
                pass
            else:
                continue
        parts.append(l)
        if len(" ".join(parts)) > 180:
            break
    if not parts:
        return None
    title = re.sub(r"\s+", " ", " ".join(parts)).strip()
    if ministry and re.match(r"^قرار\s+(رقم|وزار)", title):
        who = "وزير العمل" if "العمل" in ministry else ministry
        title = re.sub(r"^قرار\s+(وزاري|وزارى)?\s*", f"قرار {who} " if "العمل" in ministry
                       else "قرار وزاري ", title)
    title = re.sub(r"(مجلس الوزراء)\s+قرار\s+رقم", r"\1 رقم", title)
    return to_latin_digits(title)


def legal_sort_key(title: str):
    t = re.sub(r"^\d+\s+\S+\s+\d{4}\s*-\s*", "", to_latin_digits(title or ""))
    rank = 3
    if re.match(r"^(ال)?قانون", t):
        rank = 0 if "بإصدار" in t or "بشأن إصدار" in t or "إصدار" in t[:40] else 1
    elif "رئيس الجمهورية" in t or "رئيس جمهورية" in t or "مرسوم" in t:
        rank = 1
    elif "مجلس الوزراء" in t:
        rank = 2
    m = re.search(r"رقم\s*(\d+)\s*لسن[ةه]\s*(\d{4})", t)
    num, year = (int(m.group(1)), int(m.group(2))) if m else (0, 0)
    return (rank, year, num)


def strip_review_preamble(blocks):
    """Review copies start with a memo about a sample company; the regulation
    itself starts at 'ديباجة'."""
    texts = [t for _, t in blocks[:15]]
    if not any("المراجعة" in t for t in texts):
        return blocks
    for i, (_, t) in enumerate(blocks):
        if t.strip() == "ديباجة":
            return blocks[:4] + blocks[i:]
    return blocks


# ------------------------------------------------------------ forms book

def split_forms_book(docx_path: Path, did_prefix: str):
    """Splits a Word file whose forms are tables starting with 'نموذج /'.
    Returns a list of (group, title, blocks, docx_bytes_path) or [] when
    the file is not such a book."""
    import docx
    from docx.table import Table

    d = docx.Document(str(docx_path))
    body = d.element.body
    children = list(body.iterchildren())

    def table_head(el):
        rows = _table_rows(Table(el, d))
        first = " ".join(" ".join(r) for r in rows[:2])
        return rows, first

    starts = [i for i, el in enumerate(children)
              if el.tag == W_NS + "tbl" and "نموذج /" in table_head(el)[1][:300]]
    if len(starts) < 3:
        return []

    forms, group = [], "نماذج عامة"
    sect_pr = body.find(W_NS + "sectPr")
    for n, start in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(children)
        els = [el for el in children[start:end] if el.tag in (W_NS + "tbl", W_NS + "p")]
        rows, _ = table_head(els[0])
        flat = [c for r in rows[:3] for c in r]
        title = next((c for c in flat if c.startswith("نموذج /")), flat[0])
        grp = next((c for c in flat if c.startswith("المجموعة")), None)
        if grp:
            group = re.sub(r"^المجموعة\s+\S+\s*:\s*", "", grp).strip() or grp
        title = clean(title.replace("نموذج /", "").strip())

        blocks = []
        for el in els:
            if el.tag == W_NS + "tbl":
                for r in _table_rows(Table(el, d)):
                    line = clean(" | ".join(c for c in r if c not in (title, grp) and "نموذج /" not in c))
                    if line and not re.fullmatch(r"[_ ./|:]+", line):
                        blocks.append(("p", line))
        # One standalone .docx per form: same styles, only this form's content.
        single = docx.Document(str(docx_path))
        sb = single.element.body
        for el in list(sb.iterchildren()):
            sb.remove(el)
        for el in els:
            sb.append(deepcopy(el))
        if sect_pr is not None:
            sb.append(deepcopy(sect_pr))
        out = CACHE / "forms" / f"{did_prefix}_{n:03d}.docx"
        out.parent.mkdir(parents=True, exist_ok=True)
        single.save(str(out))
        forms.append((group, title, blocks, out))
    return forms


# ----------------------------------------------------------- spreadsheets

def xlsx_documents(path: Path):
    """Each data row of a table-like sheet becomes a section:
    title = short label (or start of the longest text), body = 'column: value'
    lines. Returns (title, sections); title comes from a caption cell above
    the header row when there is one."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    seen, sections, doc_title = [], [], None
    for ws in wb.worksheets:
        rows = [[clean("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 4), None)
        if hi is None:
            continue
        caption = next((c for r in rows[:hi] for c in r if c), None)
        base = [h.replace("العقوية", "العقوبة") for h in rows[hi]]
        head = base[:]
        sub = rows[hi + 1] if hi + 1 < len(rows) else []
        n_sub = sum(1 for c in sub if c)
        used_sub = 0 < n_sub <= 3
        ranges = {}  # column of "لا تقل" → column of "لا تزيد"
        if used_sub:
            for j, c in enumerate(sub):
                if c:
                    owner = next(k for k in range(j, -1, -1) if base[k])
                    head[j] = f"{base[owner]} ({c})"
                    if "تقل" in c:
                        ranges[j] = None
                    elif "تزيد" in c and ranges:
                        ranges[max(ranges)] = j
        data = rows[hi + (2 if used_sub else 1):]
        data = [r for r in data if sum(1 for c in r if c) >= 2]
        names = {h for h in head if h}
        if any(len(names & prev) >= len(names) / 2 and n == len(data) for prev, n in seen):
            continue  # a sheet that repeats an earlier one
        seen.append((names, len(data)))
        doc_title = doc_title or caption
        flags = {j for j in range(len(head))
                 if all(r[j] in ("", "0", "1") for r in data if j < len(r))
                 and any(j < len(r) and r[j] for r in data)}
        for r in data:
            lines, short, longest = [], "", ""
            for j, v in enumerate(r):
                if not v or j >= len(head) or not head[j] or j in ranges.values():
                    continue
                if j in ranges:
                    hi_v = r[ranges[j]] if ranges[j] is not None and ranges[j] < len(r) else ""
                    label = base[next(k for k in range(j, -1, -1) if base[k])]
                    lines.append(f"{label}: من {v} إلى {hi_v}" if hi_v else f"{label}: {v}")
                    continue
                if j in flags:
                    v = "نعم" if v == "1" else "لا"
                elif len(v) > len(longest):
                    longest = v
                if not short and j not in flags and 3 <= len(v) <= 60 \
                        and not re.fullmatch(r"[\d.\s]+", v) and "المادة" not in v:
                    short = v
                lines.append(f"{head[j]}: {v}")
            title = short or (longest[:70] + "…" if len(longest) > 70 else longest)
            art = next((r[j] for j, h in enumerate(head)
                        if "رقم المادة" in h and j < len(r) and r[j]), None)
            if art and art not in title:
                title = f"مادة {art}: {title}"
            sections.append({"title": title, "body": "\n".join(lines)})
    return doc_title, sections


# ------------------------------------------------------------------ main

def category_for(rel: str):
    if re.search(r"نماذج|نموذج|كشف", rel) and "لائحة نموذجية" not in rel:
        return "forms"
    if "لائحة نموذجية" in rel or "لائحه نموذجيه" in rel:
        return "work_rules"
    if re.search(r"الالتزامات|العقوبات|جزاءات", rel):
        return "obligations"
    return "laws"


def doc_id(key: str) -> str:
    return "d" + hashlib.md5(key.encode("utf-8")).hexdigest()[:10]


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="المجلد الذي يحتوي على ملفات الموارد البشرية")
    args = ap.parse_args()
    base = Path(args.folder).expanduser().resolve()
    if not base.is_dir():
        sys.exit(f"المجلد غير موجود: {base}")

    FORMS_DIR.mkdir(parents=True, exist_ok=True)
    for old in FORMS_DIR.iterdir():
        if old.is_file() and old.name != ".gitkeep":
            old.unlink()

    files = sorted(p for p in base.rglob("*")
                   if p.is_file() and not p.name.startswith(("~$", ".")) and p.stat().st_size)
    # The same model regulation in several output folders → keep the newest.
    newest = {}
    for p in files:
        if category_for(p.name) == "work_rules":
            if p.name not in newest or p.stat().st_mtime > newest[p.name].stat().st_mtime:
                newest[p.name] = p
    files = [p for p in files if category_for(p.name) != "work_rules" or newest[p.name] == p]
    stems = {(p.parent, p.stem) for p in files if p.suffix.lower() != ".pdf"}
    names = {p.stem for p in files}

    documents, skipped, seen_text = [], [], {}
    for path in files:
        rel = str(path.relative_to(base))
        ext = path.suffix.lower()
        folder = path.parent.name if path.parent != base else ""
        cat = category_for(rel)
        did = doc_id(rel)
        try:
            if ext == ".pdf" and (path.parent, path.stem) in stems:
                skipped.append(f"{rel} (مكرر: نفس الملف موجود بصيغة أخرى)")
                continue
            if cat == "work_rules" and not path.stem.endswith("نسخة المراجعة") \
                    and f"{path.stem} - نسخة المراجعة" in names:
                continue  # attached to its review copy below
            note, file_src, title, group, summary_hint = None, None, None, None, None
            if ext == ".xlsx":
                if cat == "forms":  # a sheet to fill in: show its text as-is
                    sections, file_src = to_sections(read_xlsx(path)), path
                    group = re.split(r"\s+-\s+", path.stem)[0].strip()
                else:
                    title, sections = xlsx_documents(path)
            elif ext in (".md", ".txt"):
                blocks = read_text(path)
                sections = to_sections(blocks)
                title = next((t for k, t in blocks if k == "h"), None)
            elif ext == ".pdf":
                blocks, ocr = read_pdf(path)
                if ocr:
                    note = ("نص مستخرج آلياً من ملف مصوّر (OCR) وقد يحتوي على أخطاء "
                            "في بعض الكلمات أو الأرقام. يُرجع للنسخة الرسمية عند الاستشهاد.")
                sections = to_sections(blocks)
                title = legal_title(blocks)
            elif ext in (".docx", ".rtf", ".doc"):
                src = path if ext == ".docx" else convert_to_docx(path)
                if cat == "forms":
                    forms = split_forms_book(src, did)
                    if forms:
                        for n, (grp, ftitle, blocks, fpath) in enumerate(forms):
                            fid = f"{did}_{n:03d}"
                            shutil.copy2(fpath, FORMS_DIR / f"{fid}.docx")
                            documents.append({
                                "id": fid, "category": "forms", "group": grp,
                                "title": ftitle, "summary": grp,
                                "sections": to_sections(blocks) or [{"title": "", "body": ftitle}],
                                "file": f"forms/{fid}.docx"})
                        log(f"✓ [نماذج] {rel} — {len(forms)} نموذج")
                        continue
                    file_src = src
                import docx
                blocks = docx_blocks(docx.Document(str(src)))
                if cat == "work_rules":
                    blocks = strip_review_preamble(blocks)
                    clean_copy = path.with_name(path.name.replace(" - نسخة المراجعة", ""))
                    if clean_copy != path and clean_copy.exists():
                        file_src = clean_copy
                    title = re.sub(r"\s*-\s*نسخة المراجعة$", "", path.stem)
                    note = ("لائحة نموذجية لمنشأة افتراضية. يظهر تحت كل مادة درجتها "
                            "(إلزامي / موصى به / اختياري) وسندها القانوني.")
                elif cat == "laws":
                    title = legal_title(blocks)
                sections = to_sections(blocks)
            else:
                skipped.append(f"{rel} (نوع ملف غير مدعوم)")
                continue
        except Exception as e:  # keep going; report at the end
            skipped.append(f"{rel} ({e})")
            continue

        if not sections:
            skipped.append(f"{rel} (بدون نص)")
            continue
        key = hashlib.md5("".join(s["body"] for s in sections).encode()).hexdigest()
        if key in seen_text:
            skipped.append(f"{rel} (مكرر لـ {seen_text[key]})")
            continue
        seen_text[key] = rel

        if cat == "work_rules":
            summary_hint = "لائحة نموذجية مع سند كل مادة وجدول الجزاءات"
        if cat == "laws":
            group = GROUP_NAMES.get(folder, folder) or "أخرى"
            if note and re.search(r"[\u0600-\u06ff]", path.stem):
                title = path.stem.strip()  # OCR'd headers are unreliable
            elif note and title:
                m = re.search(r"(بإصدار|بشأن|فى شأن|في شأن).*", title)
                kind = re.match(r"^\S+(\s+(وزاري|وزارى))?", title.split("رقم")[0].strip())
                title = f"{kind.group(0) if kind else ''} {m.group(0)}".strip() if m else None
            if title and not re.search(r"(بشأن|بإصدار|باصدار|بتعديل|في شأن|فى شأن|بتحديد|بإنشاء|بتشكيل)", title):
                first = next((s["body"] for s in sections if s["title"]), "")
                if first:
                    first = " ".join(first.split())
                    summary_hint = first[:110] + ("…" if len(first) > 110 else "")
        doc = {"id": did, "category": cat, "title": title or path.stem.strip(),
               "summary": summary_hint or group or "", "sections": sections}
        if group:
            doc["group"] = group
        if note:
            doc["note"] = note
        if file_src:
            name = f"{did}{file_src.suffix.lower()}"
            shutil.copy2(file_src, FORMS_DIR / name)
            doc["file"] = f"forms/{name}"
        documents.append(doc)
        log(f"✓ [{cat}] {rel} — {len(sections)} فقرة")

    # Folders named by number ("1-2017") get the title of their main law.
    for g in {d.get("group") for d in documents if d.get("group")}:
        if not re.search(r"[\u0600-\u06ff]", g):
            docs = sorted((d for d in documents if d.get("group") == g),
                          key=lambda d: legal_sort_key(d["title"]))
            nice = re.sub(r"^.*?لسن[ةه]\s*\d{4}\s*", "", docs[0]["title"])[:60] or g
            nice = re.sub(r"^(بإصدار|بشأن|في شأن|فى شأن)\s+", "", nice)
            for d in docs:
                d["group"] = d["summary"] = nice

    def order(d):
        cat_i = [c["id"] for c in CATEGORIES].index(d["category"])
        g = d.get("group") or ""
        gi = GROUP_ORDER.index(g) if g in GROUP_ORDER else len(GROUP_ORDER)
        inner = legal_sort_key(d["title"]) if d["category"] == "laws" else (0, 0, 0)
        return (cat_i, gi, g if d["category"] == "laws" else "", inner)
    documents.sort(key=order)  # stable: forms keep book order

    used = {d["category"] for d in documents}
    cats = [c for c in CATEGORIES if c["id"] in used]
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps({"categories": cats, "documents": documents},
                                   ensure_ascii=False, separators=(",", ":")),
                        encoding="utf-8")
    size = OUT_JSON.stat().st_size / 1e6
    log(f"\nتم: {len(documents)} مستند في {len(cats)} قسم → {OUT_JSON} ({size:.1f} MB)")
    if skipped:
        log("\nملفات لم يتم استيرادها:")
        for s in skipped:
            log("  - " + s)


if __name__ == "__main__":
    main()
